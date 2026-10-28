import argparse
import csv
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
from bs4 import BeautifulSoup


def inspect_article(path, url):
    source = Path(path)
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    text_path = source.with_suffix(".txt")
    text_path.write_text(soup.get_text("\n", strip=True), encoding="utf-8")
    links = [
        {"text": link.get_text(" ", strip=True), "url": urljoin(url, link["href"])}
        for link in soup.find_all("a", href=True)
        if any(term in (link.get_text(" ", strip=True) + link["href"]).lower()
               for term in ("suppl", "module", "mmc", "nihms"))
    ]
    references = []
    for ref in soup.select(".ref-list li"):
        if "Chaussabel" in ref.get_text() or "Chiche" in ref.get_text():
            references.append({"text": ref.get_text(" ", strip=True),
                               "links": [urljoin(url, link["href"])
                                         for link in ref.find_all("a", href=True)]})
    source.with_name(source.stem + "_links.json").write_text(
        json.dumps(links, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps({"text_path": str(text_path), "text_characters": text_path.stat().st_size,
                      "links": links, "references": references}, ensure_ascii=False, indent=2))


def inspect_workbook(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summary = []
    for sheet in workbook.worksheets:
        head = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= 8:
                break
            head.append(list(row[:12]))
        summary.append({"sheet": sheet.title, "rows": sheet.max_row,
                        "columns": sheet.max_column, "head": head})
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))


def source_manifest(directory):
    directory = Path(directory)
    sources = [
        ("article.html", "https://pmc.ncbi.nlm.nih.gov/articles/PMC5426482/", "Banchereau2016_full_text"),
        ("mmc1.pdf", "https://ars.els-cdn.com/content/image/1-s2.0-S0092867416302641-mmc1.pdf", "Banchereau2016_supplement"),
        ("chaussabel2014.html", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4118927/", "G2_method_and_module_mean"),
        ("tcgsa_article.html", "https://journals.plos.org/ploscompbiol/article?id=10.1371/journal.pcbi.1004310", "archived_G2_probe_membership_publication"),
        ("TcGSA_s007.zip", "https://journals.plos.org/ploscompbiol/article/file?type=supplementary&id=10.1371/journal.pcbi.1004310.s007", "complete_G2_probe_membership"),
        ("tmod_0.50.13.tar.gz", "https://cran.r-project.org/src/contrib/tmod_0.50.13.tar.gz", "complete_G2_gene_membership"),
        ("tmod_vignette.html", "https://www.stats.bris.ac.uk/R/web/packages/tmod/vignettes/tmod.html", "G2_gene_representation_documentation"),
        ("GPL10558.annot.gz", "https://ftp.ncbi.nlm.nih.gov/geo/platforms/GPL10nnn/GPL10558/annot/GPL10558.annot.gz", "official_curated_probe_annotation"),
        ("GPL10558_submitter_complete.soft", "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?targ=self&acc=GPL10558&form=text&view=full", "official_original_probe_annotation"),
        ("ncbi_gene_summary.json", "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=gene&id=129607,26010,728358&retmode=json", "official_LOC_Gene_ID_nomenclature"),
    ]
    rows = []
    for filename, url, role in sources:
        path = directory / filename
        stat = path.stat()
        if stat.st_size <= 0:
            raise ValueError("A complete public module source is empty: " + filename)
        rows.append({"file": filename, "source_url": url, "source_role": role,
                     "byte_size": stat.st_size,
                     "file_completed_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                     "local_path": str(path.resolve()), "status": "complete"})
    with (directory / "source_manifest.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(json.dumps({"source_files": len(rows), "manifest": str(directory / "source_manifest.csv")},
                     ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--article", required=True)
    inspect_parser.add_argument("--url", required=True)
    workbook_parser = subparsers.add_parser("inspect-workbook")
    workbook_parser.add_argument("--workbook", required=True)
    manifest_parser = subparsers.add_parser("source-manifest")
    manifest_parser.add_argument("--directory", required=True)
    args = parser.parse_args()
    if args.command == "inspect":
        inspect_article(args.article, args.url)
    elif args.command == "inspect-workbook":
        inspect_workbook(args.workbook)
    elif args.command == "source-manifest":
        source_manifest(args.directory)


if __name__ == "__main__":
    main()
